#!/usr/bin/env python3
"""Create Excel starter workbooks for metrics / AOI tables."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

OUT = Path(__file__).resolve().parent.parent / "excel"
HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HINT_FILL = PatternFill("solid", fgColor="F4EDE4")


def style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.font = HEADER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 18


def make_metrics_starter() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "FixationLog"
    headers = [
        "participant",
        "trial",
        "fixation_id",
        "start_s",
        "end_s",
        "duration_s",
        "aoi",
    ]
    ws.append(headers)
    # Sample rows (students replace with exports)
    samples = [
        ("p01", 1, 1, 0.00, 0.32, None, "logo"),
        ("p01", 1, 2, 0.45, 0.70, None, "price"),
        ("p01", 1, 3, 0.90, 1.40, None, "buy_button"),
        ("p02", 1, 1, 0.00, 0.50, None, "logo"),
        ("p02", 1, 2, 0.60, 0.85, None, "price"),
    ]
    for row in samples:
        ws.append(list(row))
    # duration formula
    for r in range(2, 2 + len(samples)):
        ws.cell(r, 6, f"=E{r}-D{r}")
    style_header(ws, len(headers))

    ws2 = wb.create_sheet("SummaryFormulas")
    ws2["A1"] = "Metric"
    ws2["B1"] = "Excel formula (edit ranges as needed)"
    ws2["C1"] = "Result (demo)"
    rows = [
        ("Total dwell (s)", "=SUM(FixationLog!F:F)", "=SUM(FixationLog!F2:F6)"),
        ("N fixations", "=COUNTA(FixationLog!C2:C100)", None),
        ("Mean fixation (s)", "=AVERAGE(FixationLog!F2:F100)", None),
        ("Dwell on logo", '=SUMIF(FixationLog!G:G,"logo",FixationLog!F:F)', None),
    ]
    for i, (m, f, _) in enumerate(rows, start=2):
        ws2.cell(i, 1, m)
        ws2.cell(i, 2, f)
        if m == "Total dwell (s)":
            ws2.cell(i, 3, "=SUM(FixationLog!F2:F6)")
        elif m == "N fixations":
            ws2.cell(i, 3, '=COUNTA(FixationLog!C2:C6)')
        elif m == "Mean fixation (s)":
            ws2.cell(i, 3, "=AVERAGE(FixationLog!F2:F6)")
        else:
            ws2.cell(i, 3, '=SUMIF(FixationLog!G2:G6,"logo",FixationLog!F2:F6)')
    style_header(ws2, 3)
    ws2.column_dimensions["B"].width = 48

    ws3 = wb.create_sheet("HowTo")
    ws3["A1"] = "How to use this workbook"
    ws3["A1"].font = Font(bold=True, size=14)
    tips = [
        "1. Paste fixation exports into FixationLog (or import CSV).",
        "2. Keep duration_s as a formula (end - start) when possible.",
        "3. Use SummaryFormulas as a template for your own lab sheet.",
        "4. Pair with Python notebooks for larger N and plots.",
        "5. Always keep a raw copy of exports untouched.",
    ]
    for i, t in enumerate(tips, start=3):
        ws3.cell(i, 1, t)
        ws3.cell(i, 1).fill = HINT_FILL
    ws3.column_dimensions["A"].width = 80

    wb.save(OUT / "metrics-starter.xlsx")


def make_aoi_starter() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "AOI_Metrics"
    headers = [
        "Participant",
        "Media",
        "AOI",
        "Number_of_fixations",
        "Total_duration_of_fixations",
        "Time_to_first_fixation",
        "condition",
    ]
    ws.append(headers)
    demo = [
        ("p01", "cake", "product", 8, 1200, 180, "high_cal"),
        ("p01", "cake", "buy_button", 2, 300, 900, "high_cal"),
        ("p02", "cake", "product", 5, 800, 220, "low_cal"),
        ("p02", "cake", "buy_button", 4, 700, 400, "low_cal"),
        ("p03", "cereal", "product", 6, 950, 150, "high_cal"),
        ("p03", "cereal", "buy_button", 1, 120, 1100, "high_cal"),
    ]
    for row in demo:
        ws.append(list(row))
    style_header(ws, len(headers))

    pivot = wb.create_sheet("Pivot_Dwell_by_AOI")
    pivot.append(["AOI", "Mean_dwell_ms", "Mean_TTFF_ms"])
    pivot.append(["product", "=AVERAGEIF(AOI_Metrics!C:C,A2,AOI_Metrics!E:E)", "=AVERAGEIF(AOI_Metrics!C:C,A2,AOI_Metrics!F:F)"])
    pivot.append(["buy_button", "=AVERAGEIF(AOI_Metrics!C:C,A3,AOI_Metrics!E:E)", "=AVERAGEIF(AOI_Metrics!C:C,A3,AOI_Metrics!F:F)"])
    style_header(pivot, 3)

    # Precomputed values for environments that don't recalc AVERAGEIF across columns the same way
    values = wb.create_sheet("QuickChartData")
    values.append(["AOI", "Mean_dwell_ms"])
    values.append(["product", 983.3])
    values.append(["buy_button", 373.3])
    style_header(values, 2)
    chart = BarChart()
    chart.title = "Mean dwell by AOI (demo numbers)"
    chart.y_axis.title = "ms"
    data = Reference(values, min_col=2, min_row=1, max_row=3)
    cats = Reference(values, min_col=1, min_row=2, max_row=3)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    values.add_chart(chart, "D2")

    howto = wb.create_sheet("HowTo")
    howto["A1"] = "Import Data/demo/tobii_like_aoi_metrics.csv into AOI_Metrics (replace sample rows)."
    howto["A2"] = "Rebuild pivots: Insert → PivotTable (Excel) or Data → Pivot (LibreOffice)."
    howto["A3"] = "TTFF = Time to first fixation; lower often means earlier attention (context matters)."
    howto.column_dimensions["A"].width = 90

    wb.save(OUT / "aoi-tables-starter.xlsx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_metrics_starter()
    make_aoi_starter()
    print(f"Wrote Excel starters to {OUT}")


if __name__ == "__main__":
    main()
